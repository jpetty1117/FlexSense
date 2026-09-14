from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor

FONT_NAME = "Inter"
BLACK = RGBColor(0x00, 0x00, 0x00)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
HEADER_FILL = RGBColor(0x22, 0x22, 0x22)
BODY_FILL = RGBColor(0x00, 0x00, 0x00)

SLIDE_W = Inches(13.333)
SLIDE_H = Inches(7.5)

SECTION_ORDER = [
    ("accomplishment", "Accomplishments / Completed Goals"),
    ("challenge", "Challenges / Missed Goals / Help Needed"),
    ("goal", "Goals for Next Week"),
]


def _new_presentation():
    prs = Presentation()
    prs.slide_width = SLIDE_W
    prs.slide_height = SLIDE_H
    return prs


def _blank_slide(prs):
    layout = prs.slide_layouts[6]  # Blank
    slide = prs.slides.add_slide(layout)
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = BLACK
    return slide


def _add_title(slide, text, subtitle=None):
    box = slide.shapes.add_textbox(Inches(0.6), Inches(0.35), SLIDE_W - Inches(1.2), Inches(1.0))
    tf = box.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    run = p.add_run()
    run.text = text
    run.font.size = Pt(32)
    run.font.bold = True
    run.font.name = FONT_NAME
    run.font.color.rgb = WHITE
    if subtitle:
        sp = tf.add_paragraph()
        srun = sp.add_run()
        srun.text = subtitle
        srun.font.size = Pt(15)
        srun.font.name = FONT_NAME
        srun.font.color.rgb = WHITE


def fmt_num(v):
    return f"{v:.2f}"


def _style_cell(cell, bold, fill):
    cell.fill.solid()
    cell.fill.fore_color.rgb = fill
    cell.margin_left = Pt(6)
    cell.margin_right = Pt(6)
    tf = cell.text_frame
    for p in tf.paragraphs:
        for run in p.runs:
            run.font.size = Pt(13)
            run.font.bold = bold
            run.font.name = FONT_NAME
            run.font.color.rgb = WHITE


def build_timesheet_slide(prs, member_rows, team_total, remaining_hours_val, weeks_left_val, pace_val, subtitle=None):
    slide = _blank_slide(prs)
    _add_title(slide, "Labor Status", subtitle=subtitle)

    n_rows = len(member_rows) + 1 + (1 if team_total else 0)
    n_cols = 6
    left, top = Inches(0.6), Inches(1.5)
    width, height = Inches(10.6), Inches(0.45) * n_rows
    table = slide.shapes.add_table(n_rows, n_cols, left, top, width, height).table

    headers = ["Member", "Predicted", "Actual", "Delta", "Next Week", "Cumulative"]
    for c, h in enumerate(headers):
        table.cell(0, c).text = h
        _style_cell(table.cell(0, c), bold=True, fill=HEADER_FILL)

    r = 1
    for row in member_rows:
        vals = [row["name"], fmt_num(row["predicted"]), fmt_num(row["actual"]),
                fmt_num(row["delta"]), fmt_num(row["next_week"]), fmt_num(row["cumulative"])]
        for c, v in enumerate(vals):
            table.cell(r, c).text = v
            _style_cell(table.cell(r, c), bold=False, fill=BODY_FILL)
        r += 1

    if team_total:
        vals = ["Team Total", fmt_num(team_total["predicted"]), fmt_num(team_total["actual"]),
                fmt_num(team_total["delta"]), fmt_num(team_total["next_week"]), fmt_num(team_total["cumulative"])]
        for c, v in enumerate(vals):
            table.cell(r, c).text = v
            _style_cell(table.cell(r, c), bold=True, fill=HEADER_FILL)

    summary_top = top + height + Inches(0.35)
    box = slide.shapes.add_textbox(left, summary_top, width, Inches(1.3))
    tf = box.text_frame
    tf.word_wrap = True
    lines = [
        f"Estimated Remaining Labor: {remaining_hours_val:.0f} hours",
        f"Weeks Remaining: {weeks_left_val:.0f}",
        f"Minimum Hours / Week: {pace_val:.0f} hours",
    ]
    for i, line in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        run = p.add_run()
        run.text = line
        run.font.size = Pt(16)
        run.font.name = FONT_NAME
        run.font.color.rgb = WHITE
    return slide


def build_individual_slide(prs, full_name, role, sections):
    slide = _blank_slide(prs)
    _add_title(slide, full_name, subtitle=role)

    box = slide.shapes.add_textbox(Inches(0.6), Inches(1.6), SLIDE_W - Inches(1.2), SLIDE_H - Inches(2.0))
    tf = box.text_frame
    tf.word_wrap = True
    first = True
    for key, label in SECTION_ORDER:
        p = tf.paragraphs[0] if first else tf.add_paragraph()
        first = False
        p.space_before = Pt(0 if key == "accomplishment" else 14)
        run = p.add_run()
        run.text = label
        run.font.size = Pt(19)
        run.font.bold = True
        run.font.name = FONT_NAME
        run.font.color.rgb = WHITE

        bullets = sections.get(key) or ["No updates logged this period"]
        for b in bullets:
            bp = tf.add_paragraph()
            bp.space_before = Pt(2)
            run = bp.add_run()
            run.text = f"\u2022  {b}"
            run.font.size = Pt(14)
            run.font.name = FONT_NAME
            run.font.color.rgb = WHITE
    return slide


def generate_pptx(output_path, member_rows, team_total, remaining_hours_val, weeks_left_val, pace_val,
                   owner_sections, subtitle=None):
    """owner_sections: ordered list of dicts: {full_name, role, accomplishment, challenge, goal}"""
    prs = _new_presentation()
    build_timesheet_slide(prs, member_rows, team_total, remaining_hours_val, weeks_left_val, pace_val,
                           subtitle=subtitle)
    for data in owner_sections:
        build_individual_slide(prs, data["full_name"], data["role"],
                                {"accomplishment": data.get("accomplishment", []),
                                 "challenge": data.get("challenge", []),
                                 "goal": data.get("goal", [])})
    prs.save(output_path)
    return output_path
