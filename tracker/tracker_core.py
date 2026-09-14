"""Shared data and business logic for the Squish Therapy project tracker web app.

The baseline task plan and calculations are ported from tracker_gui.py so the
web app uses the same task IDs, dates, status semantics, timesheet calculations,
and Gantt export behavior as the desktop tracker.
"""

from __future__ import annotations

import csv
import io
import json
import os
import re
import threading
import xml.etree.ElementTree as ET
from datetime import date, datetime, timedelta
from pathlib import Path
from xml.dom import minidom

try:
    import openpyxl
except ImportError:  # handled gracefully by the Flask route
    openpyxl = None

BASE_DIR = Path(__file__).resolve().parent
DATA_DIR = BASE_DIR / "data"
DATA_DIR.mkdir(exist_ok=True)
STATE_FILE = DATA_DIR / "tracker_state.json"
TIMESHEET_STATE_FILE = DATA_DIR / "timesheet_state.json"
APP_META_FILE = DATA_DIR / "tracker_meta.json"
BOM_STATE_FILE = DATA_DIR / "bom_state.json"
CAD_STATE_FILE = DATA_DIR / "cad_state.json"
CAD_FILES_DIR = DATA_DIR / "cad_models"
CAD_FILES_DIR.mkdir(exist_ok=True)

STATUSES = ['not-started', 'in-progress', 'complete', 'blocked']

HOURS_PER_TASK_DAY = 2

OWNER_KEYS = ['Cyrus', 'Mya', 'Ana', 'Nick', 'Jase']

OWNER_INFO = {'Cyrus': ('Cyrus Lowden', 'Project Manager'),
 'Mya': ('Mya Tinsay', 'Hardware Engineer'),
 'Ana': ('Ana Luo', 'Mechanical Engineer'),
 'Nick': ('Nick Bailey', 'Test Engineer'),
 'Jase': ('Jase Petty', 'Software Engineer')}

NOTE_CATEGORIES = [('accomplishment', 'Accomplishment'),
 ('challenge', 'Challenge / Help Needed'),
 ('goal', 'Goal for Next Week'),
 ('note', 'General Note')]
NOTE_CATEGORY_LABELS = dict(NOTE_CATEGORIES)


STATUS_COMPLETION = {'not-started': 0, 'in-progress': 50, 'blocked': 25, 'complete': 100}

GAN_COLORS = {'not-started': '#B9C4CC', 'in-progress': '#2B5A7A', 'blocked': '#9C4430', 'complete': '#3E7A5C'}

SORT_OPTIONS = ['Plan order', 'Due date (soonest first)', 'Most overdue first', 'Name (A–Z)', 'Status', 'Owner (A–Z)']

DATA = [('1.1.1',
  '1. Electronics',
  '1.1 Microcontroller Unit (MCU)',
  'Research and choose affordable MCU',
  '3/23/26',
  '3/28/26'),
 ('1.1.2',
  '1. Electronics',
  '1.1 Microcontroller Unit (MCU)',
  'Design circuit for MCU to communicate with computer and sensors',
  '3/31/26',
  '4/7/26'),
 ('1.1.3',
  '1. Electronics',
  '1.1 Microcontroller Unit (MCU)',
  'Research MCU-based GUI applications for displaying data',
  '3/31/26',
  '4/7/26'),
 ('1.1.4',
  '1. Electronics',
  '1.1 Microcontroller Unit (MCU)',
  'Research MCU-based database implementations (Library/DB)',
  '4/9/26',
  '4/16/26'),
 ('1.1.5', '1. Electronics', '1.1 Microcontroller Unit (MCU)', 'Order and Receive MCU', '4/16/26', '4/23/26'),
 ('1.1.6',
  '1. Electronics',
  '1.1 Microcontroller Unit (MCU)',
  'Order and Receive auxiliary parts (Power, analog devices, etc.)',
  '8/26/26',
  '9/2/26'),
 ('1.1.7', '1. Electronics', '1.1 Microcontroller Unit (MCU)', 'Build circuit for MCU and Sensors', '9/2/26', '9/9/26'),
 ('1.1.8',
  '1. Electronics',
  '1.1 Microcontroller Unit (MCU)',
  'Design high-level page design for info visualization/storage',
  '9/9/26',
  '9/12/26'),
 ('1.2.1',
  '1. Electronics',
  '1.2 Torque Generation/Control System',
  'Research and Choose Force Measurement Sensor',
  '3/16/26',
  '3/19/26'),
 ('1.2.2',
  '1. Electronics',
  '1.2 Torque Generation/Control System',
  'Research for Torque Resistance Mechanism with Enough Granular Control in Desired Price Range',
  '3/19/26',
  '3/26/26'),
 ('1.2.3',
  '1. Electronics',
  '1.2 Torque Generation/Control System',
  "Research chosen sensors' outputs and communication protocols",
  '4/8/26',
  '4/15/26'),
 ('1.2.4',
  '1. Electronics',
  '1.2 Torque Generation/Control System',
  'Order and receive resistance control mechanism',
  '8/25/26',
  '9/1/26'),
 ('1.3.1',
  '1. Electronics',
  '1.3 Position/Biometric Sensor System',
  'Research and Choose Angle Measurement Sensing',
  '3/17/26',
  '3/20/26'),
 ('1.3.2',
  '1. Electronics',
  '1.3 Position/Biometric Sensor System',
  'Research and Choose Linear Distance Sensor',
  '3/20/26',
  '3/25/26'),
 ('1.3.3', '1. Electronics', '1.3 Position/Biometric Sensor System', 'Order and Receive Sensors', '8/25/26', '9/1/26'),
 ('2.1.1',
  '2. Mechanical Assembly',
  '2.1 Elbow Pad',
  'Research at Least 3 Materials for Elbow Pad',
  '3/30/26',
  '4/4/26'),
 ('2.1.2', '2. Mechanical Assembly', '2.1 Elbow Pad', 'Order mechanical aspect and fasteners', '8/31/26', '9/5/26'),
 ('2.1.3',
  '2. Mechanical Assembly',
  '2.1 Elbow Pad',
  'Assemble supporting base pad to fit rails and resistance securely',
  '9/7/26',
  '9/12/26'),
 ('2.1.4',
  '2. Mechanical Assembly',
  '2.1 Elbow Pad',
  'Test base and rails for stabilization and comfort',
  '9/14/26',
  '9/17/26'),
 ('2.2.1',
  '2. Mechanical Assembly',
  '2.2 Telescoping Rails',
  'Research Rails that slide Easily and Extend to Desired Length 9-14 Inches',
  '3/23/26',
  '3/26/26'),
 ('2.2.2',
  '2. Mechanical Assembly',
  '2.2 Telescoping Rails',
  'Design Hinge/Attachment for rails and motor resistance',
  '4/13/26',
  '4/18/26'),
 ('2.2.6',
  '2. Mechanical Assembly',
  '2.2 Telescoping Rails',
  'Adjust safety limit features (ROM of hinge)',
  '9/25/26',
  '9/30/26'),
 ('2.2.7',
  '2. Mechanical Assembly',
  '2.2 Telescoping Rails',
  'Test mechanical assembly functionality for varying arm lengths, angles, and possible load conditions',
  '11/13/26',
  '11/18/26'),
 ('2.3.1',
  '2. Mechanical Assembly',
  '2.3 Handle Grip',
  'Research handle grip and form for suitable force measurement installment and lock method',
  '3/30/26',
  '4/2/26'),
 ('2.3.2',
  '2. Mechanical Assembly',
  '2.3 Handle Grip',
  'Design fitting and attachment for handle grip, force measurement device, and rails',
  '4/23/26',
  '4/30/26'),
 ('2.3.4',
  '2. Mechanical Assembly',
  '2.3 Handle Grip',
  'Install handle grip and test for sturdability and stabilization',
  '10/6/26',
  '10/9/26'),
 ('3.1.1',
  '3. Software/Data Processing',
  '3.1 Data Acquisition',
  'Write data acquisition software for reading data from sensors based on high-level design',
  '9/14/26',
  '9/22/26'),
 ('3.1.2',
  '3. Software/Data Processing',
  '3.1 Data Acquisition',
  'Implement necessary software filters to smooth data',
  '10/1/26',
  '10/8/26'),
 ('3.1.3',
  '3. Software/Data Processing',
  '3.1 Data Acquisition',
  'Calibrate sensor outputs to useable metrics for GUI and Control',
  '10/13/26',
  '10/20/26'),
 ('3.2.1',
  '3. Software/Data Processing',
  '3.2 Control',
  'Test sensor I/O under different conditions/speeds to see noise levels',
  '9/28/26',
  '10/1/26'),
 ('3.2.2',
  '3. Software/Data Processing',
  '3.2 Control',
  'Test resistance mechanism to create a forward lookup table for control input vs resistance',
  '10/1/26',
  '10/8/26'),
 ('3.2.3',
  '3. Software/Data Processing',
  '3.2 Control',
  'Design Feedforward/Feedback loop to control torque resistance mechanism',
  '10/26/26',
  '10/31/26'),
 ('3.2.4',
  '3. Software/Data Processing',
  '3.2 Control',
  'Tune Feedback and test control system against disturbances',
  '11/9/26',
  '11/12/26'),
 ('3.2.5',
  '3. Software/Data Processing',
  '3.2 Control',
  'Implement RTOS handling for GUI, Sensor Reading and Control Logic',
  '11/16/26',
  '11/19/26'),
 ('3.3.1',
  '3. Software/Data Processing',
  '3.3 GUI',
  'Build GUI front-end with plots and navigation',
  '9/21/26',
  '9/26/26'),
 ('3.3.2',
  '3. Software/Data Processing',
  '3.3 GUI',
  'Build GUI back-end databases for front-end connection',
  '9/28/26',
  '10/3/26'),
 ('3.3.3', '3. Software/Data Processing', '3.3 GUI', 'Connect GUI to Sensor Data', '10/14/26', '10/21/26'),
 ('3.3.4',
  '3. Software/Data Processing',
  '3.3 GUI',
  'Debug GUI to make sure data can be properly displayed/retrieved',
  '10/29/26',
  '11/3/26'),
 ('3.3.5', '3. Software/Data Processing', '3.3 GUI', 'Connect MCU-GUI to show on connected PC', '11/9/26', '11/12/26'),
 ('4.0.0',
  '4. Integration',
  '4.0 System Integration',
  'Combine Mechanical Assembly, Electrical Components, and MCU Software',
  '11/19/26',
  '11/28/26')]


COLORS = {
    "blue": "#2B5A7A",
    "upcoming": "#2B5A7A",
}

BUCKET_LABEL = {
    "overdue": "OVERDUE",
    "soon": "DUE SOON",
    "upcoming": "UPCOMING",
    "complete": "COMPLETE",
}

_LOCK = threading.RLock()


def parse_mdy(s: str) -> date:
    return datetime.strptime(s, "%m/%d/%y").date()


def fmt_date(d: date) -> str:
    return f"{d.month}/{d.day}/{str(d.year)[2:]}"


TASKS = {}
for tid, sec, sub, name, start, end in DATA:
    TASKS[tid] = {
        "id": tid,
        "section": sec,
        "subsection": sub,
        "name": name,
        "start": parse_mdy(start),
        "end": parse_mdy(end),
    }
TASK_ORDER = [t[0] for t in DATA]


def _read_json(path: Path, default):
    with _LOCK:
        if path.exists():
            try:
                return json.loads(path.read_text(encoding="utf-8"))
            except (json.JSONDecodeError, OSError):
                return default
        return default


def _write_json(path: Path, payload) -> None:
    with _LOCK:
        tmp = path.with_suffix(path.suffix + ".tmp")
        tmp.write_text(json.dumps(payload, indent=2, sort_keys=True), encoding="utf-8")
        os.replace(tmp, path)


def load_state():
    return _read_json(STATE_FILE, {})


def save_state(state):
    _write_json(STATE_FILE, state)


def load_bom_state():
    """Load the editable Bill of Materials state.

    The initial file is seeded from the team's uploaded Parts List.xlsx.
    Keeping BOM data in JSON makes edits persist locally just like task and
    timesheet state.
    """
    default = {"budget": 800, "source_note": "", "columns": {"extra_label": "Weight / Extra"}, "items": []}
    data = _read_json(BOM_STATE_FILE, default)
    if not isinstance(data, dict):
        data = default.copy()
    data.setdefault("budget", 800)
    data.setdefault("source_note", "")
    data.setdefault("columns", {"extra_label": "Weight / Extra"})
    data.setdefault("items", [])
    if not isinstance(data["items"], list):
        data["items"] = []
    # Make sure every row has the fields used by the browser UI.
    fields = ("category", "part", "part_number", "vendor", "price_link", "link", "notes", "extra")
    next_id = 1
    for item in data["items"]:
        if not isinstance(item, dict):
            continue
        try:
            next_id = max(next_id, int(item.get("id", 0)) + 1)
        except (TypeError, ValueError):
            pass
        for field in fields:
            item.setdefault(field, "")
        # Backward compatibility with the previous BOM UI, which combined
        # price and purchase URL in one field. Move URL-only values into the
        # dedicated link field without changing numeric/FREE prices.
        price_value = item.get("price_link", "")
        if (not item.get("link") and isinstance(price_value, str)
                and price_value.strip().lower().startswith(("http://", "https://"))):
            item["link"] = price_value.strip()
            item["price_link"] = ""
    used_ids = set()
    for item in data["items"]:
        if not isinstance(item, dict):
            continue
        try:
            item_id = int(item.get("id"))
        except (TypeError, ValueError):
            item_id = 0
        if item_id <= 0 or item_id in used_ids:
            while next_id in used_ids:
                next_id += 1
            item_id = next_id
            next_id += 1
            item["id"] = item_id
        used_ids.add(item_id)
    return data


def save_bom_state(data):
    _write_json(BOM_STATE_FILE, data)



def load_cad_state():
    """Load CAD design/version metadata. STEP files live in data/cad_models/."""
    default = {"designs": []}
    data = _read_json(CAD_STATE_FILE, default)
    if not isinstance(data, dict):
        data = {"designs": []}
    data.setdefault("designs", [])
    if not isinstance(data["designs"], list):
        data["designs"] = []
    for design in data["designs"]:
        if not isinstance(design, dict):
            continue
        design.setdefault("id", "")
        design.setdefault("name", "Untitled Design")
        design.setdefault("created_at", "")
        design.setdefault("versions", [])
        if not isinstance(design["versions"], list):
            design["versions"] = []
        for version in design["versions"]:
            if not isinstance(version, dict):
                continue
            version.setdefault("id", "")
            version.setdefault("version_number", 0)
            version.setdefault("label", "")
            version.setdefault("original_filename", "")
            version.setdefault("stored_filename", "")
            version.setdefault("uploaded_at", "")
            version.setdefault("change_summary", "")
            version.setdefault("notes", "")
            version.setdefault("file_size", 0)
    return data


def save_cad_state(data):
    _write_json(CAD_STATE_FILE, data)


def find_cad_design(data, design_id):
    return next((d for d in data.get("designs", []) if d.get("id") == design_id), None)


def find_cad_version(design, version_id):
    if not design:
        return None
    return next((v for v in design.get("versions", []) if v.get("id") == version_id), None)


def next_cad_version_number(design):
    values = []
    for version in design.get("versions", []):
        try:
            values.append(int(version.get("version_number", 0)))
        except (TypeError, ValueError):
            pass
    return max(values, default=0) + 1

def _numeric_bom_price(value):
    if isinstance(value, (int, float)) and not isinstance(value, bool):
        return float(value)
    if value is None:
        return None
    text = str(value).strip()
    if not text or text.upper() == "FREE" or text.lower().startswith(("http://", "https://")):
        return None
    cleaned = text.replace("$", "").replace(",", "").strip()
    try:
        return float(cleaned)
    except ValueError:
        return None


def bom_summary(data):
    items = [x for x in data.get("items", []) if isinstance(x, dict)]
    numeric_prices = [_numeric_bom_price(x.get("price_link")) for x in items]
    known_cost = round(sum(v for v in numeric_prices if v is not None), 2)
    try:
        budget = float(data.get("budget", 800))
    except (TypeError, ValueError):
        budget = 800.0
    categories = {str(x.get("category", "")).strip() for x in items if str(x.get("category", "")).strip()}
    return {
        "item_count": len(items),
        "category_count": len(categories),
        "priced_count": sum(v is not None for v in numeric_prices),
        "known_cost": known_cost,
        "budget": budget,
        "remaining": round(budget - known_cost, 2),
    }


def next_bom_id(data):
    ids = []
    for item in data.get("items", []):
        try:
            ids.append(int(item.get("id", 0)))
        except (TypeError, ValueError, AttributeError):
            pass
    return (max(ids) + 1) if ids else 1


def _parse_any_date(v):
    if isinstance(v, datetime):
        return v.date()
    if isinstance(v, date):
        return v
    if v is None:
        return None
    s = str(v).strip()
    for fmt in ("%m/%d/%Y", "%m/%d/%y", "%Y-%m-%d", "%d-%b", "%b-%d", "%B %d", "%m/%d"):
        try:
            d = datetime.strptime(s, fmt)
            if fmt in ("%d-%b", "%b-%d", "%B %d", "%m/%d"):
                d = d.replace(year=date.today().year)
            return d.date()
        except ValueError:
            continue
    return None


def _migrate_entry(entry):
    if "note_entries" not in entry:
        old = entry.get("notes", "") or ""
        entries = []
        for line in old.split("\n"):
            line = line.strip()
            if not line:
                continue
            m = re.search(r"-\s*(\d{1,2}/\d{1,2}(?:/\d{2})?)\s*Update\s*$", line, re.IGNORECASE)
            if m:
                date_str = m.group(1)
                text = line[:m.start()].strip(" -")
                d = _parse_any_date(date_str)
            else:
                text = line
                d = None
            entries.append({
                "date": d.isoformat() if d else "",
                "category": "note",
                "text": text,
            })
        entry["note_entries"] = entries
    entry.setdefault("owner", "")
    entry.setdefault("status", "not-started")
    entry.setdefault("revised", "")
    entry.setdefault("notes", "")
    return entry


def get_entry(state, tid):
    entry = state.setdefault(tid, {
        "status": "not-started",
        "revised": "",
        "notes": "",
        "owner": "",
        "note_entries": [],
    })
    return _migrate_entry(entry)


def normalize_state(state):
    for tid in TASK_ORDER:
        get_entry(state, tid)
    return state


def bucket_for(task, entry, today=None):
    today = today or date.today()
    if entry["status"] == "complete":
        return "complete"
    days = (task["end"] - today).days
    if days < 0:
        return "overdue"
    if days <= 14:
        return "soon"
    return "upcoming"


def remaining_hours(state_data):
    total = 0
    for tid in TASK_ORDER:
        entry = get_entry(state_data, tid)
        if entry["status"] == "complete":
            continue
        task = TASKS[tid]
        days = max(1, (task["end"] - task["start"]).days)
        total += days * HOURS_PER_TASK_DAY
    return total


def load_or_init_demo_date():
    meta = _read_json(APP_META_FILE, {})
    try:
        return datetime.strptime(meta["demo_date"], "%Y-%m-%d").date()
    except (KeyError, ValueError, TypeError):
        d = date.today() + timedelta(weeks=12)
        meta["demo_date"] = d.isoformat()
        _write_json(APP_META_FILE, meta)
        return d


def load_meta_value(key, default=None):
    return _read_json(APP_META_FILE, {}).get(key, default)


def save_meta(key, value):
    meta = _read_json(APP_META_FILE, {})
    meta[key] = value
    _write_json(APP_META_FILE, meta)


def load_timesheet_state():
    return _read_json(TIMESHEET_STATE_FILE, {})


def save_timesheet_state(state):
    _write_json(TIMESHEET_STATE_FILE, state)


def _to_float(v):
    try:
        return float(v)
    except (TypeError, ValueError):
        return 0.0


def _find_label_value(ws, label_text):
    label_norm = label_text.strip().lower()
    for row in ws.iter_rows():
        for cell in row:
            if cell.value is not None and str(cell.value).strip().lower() == label_norm:
                r = cell.row
                for c in range(cell.column + 1, ws.max_column + 1):
                    v = ws.cell(row=r, column=c).value
                    if v is not None and str(v).strip() != "":
                        return v
                return None
    return None


def parse_timesheet(file_or_path):
    if openpyxl is None:
        raise RuntimeError("openpyxl is not installed")
    wb = openpyxl.load_workbook(file_or_path, data_only=True)
    ws = wb.active

    name = _find_label_value(ws, "Name:")
    if not name:
        raise ValueError("could not find a 'Name:' label")

    predicted_this = _to_float(_find_label_value(ws, "Predicted Hours for this Week:"))
    predicted_next = _to_float(_find_label_value(ws, "Predicted Hours for next Week:"))

    header_row = None
    col_map = {}
    for row in ws.iter_rows():
        values = {str(c.value).strip().lower(): c.column for c in row if c.value}
        if "date" in values and "activity description" in values:
            header_row = row[0].row
            col_map = values
            break
    if header_row is None:
        raise ValueError("could not find the timesheet table header row")

    date_col = col_map.get("date")
    total_col = col_map.get("total time")
    desc_col = col_map.get("activity description")

    actual = 0.0
    latest_date = None
    for r in range(header_row + 1, ws.max_row + 1):
        desc = ws.cell(row=r, column=desc_col).value if desc_col else None
        if desc is None or str(desc).strip() == "":
            continue
        actual += _to_float(ws.cell(row=r, column=total_col).value if total_col else None)
        parsed = _parse_any_date(ws.cell(row=r, column=date_col).value if date_col else None)
        if parsed and (latest_date is None or parsed > latest_date):
            latest_date = parsed

    if latest_date is None:
        latest_date = date.today()

    return {
        "name": str(name).strip(),
        "predicted": round(predicted_this, 2),
        "next_week": round(predicted_next, 2),
        "actual": round(actual, 2),
        "week_key": latest_date.isoformat(),
    }


def compute_labor_status_rows(timesheet_state):
    member_rows = []
    totals = {"predicted": 0.0, "actual": 0.0, "delta": 0.0, "next_week": 0.0, "cumulative": 0.0}
    for name in sorted(timesheet_state.keys()):
        weeks = timesheet_state[name]
        if not weeks:
            continue
        latest_key = max(weeks.keys())
        latest = weeks[latest_key]
        predicted = latest["predicted"]
        actual = latest["actual"]
        next_week = latest["next_week"]
        delta = round(actual - predicted, 2)
        cumulative = round(sum(w["actual"] for w in weeks.values()), 2)
        member_rows.append({"name": name, "predicted": predicted, "actual": actual,
                            "delta": delta, "next_week": next_week, "cumulative": cumulative})
        totals["predicted"] += predicted
        totals["actual"] += actual
        totals["delta"] += delta
        totals["next_week"] += next_week
        totals["cumulative"] += cumulative
    if not member_rows:
        return member_rows, None
    for k in totals:
        totals[k] = round(totals[k], 2)
    return member_rows, totals


def compute_week_number(meeting_date, demo_date, weeks_total=12):
    project_start = demo_date - timedelta(weeks=weeks_total)
    week = ((meeting_date - project_start).days // 7) + 1
    return max(1, week)


def leaf_end(state, tid):
    entry = get_entry(state, tid)
    revised = entry.get("revised", "").strip()
    if revised:
        try:
            return datetime.strptime(revised, "%Y-%m-%d").date()
        except ValueError:
            pass
    return TASKS[tid]["end"]


def build_csv_bytes(state, today=None):
    today = today or date.today()
    buf = io.StringIO()
    writer = csv.writer(buf)
    writer.writerow(["ID", "Section", "Subsection", "Name", "Planned Start", "Planned End",
                     "Revised End", "Status", "Bucket", "Owner", "Notes"])
    for tid in TASK_ORDER:
        task = TASKS[tid]
        entry = get_entry(state, tid)
        bucket = bucket_for(task, entry, today)
        note_text = " | ".join(
            f"[{n.get('date') or 'undated'}] {NOTE_CATEGORY_LABELS.get(n.get('category'), n.get('category'))}: {n.get('text','')}"
            for n in entry.get("note_entries", [])
        )
        writer.writerow([
            task["id"], task["section"], task["subsection"], task["name"],
            fmt_date(task["start"]), fmt_date(task["end"]), entry.get("revised", ""),
            entry["status"], BUCKET_LABEL[bucket], entry.get("owner", ""), note_text,
        ])
    return buf.getvalue().encode("utf-8-sig")


def build_gan_bytes(state, today=None):
    today = today or date.today()
    sections = {}
    for tid in TASK_ORDER:
        task = TASKS[tid]
        sections.setdefault(task["section"], {}).setdefault(task["subsection"], []).append(tid)

    root = ET.Element("project", {
        "name": "Project Tracker", "company": "", "webLink": "", "view-date": today.isoformat(),
        "view-index": "0", "gantt-divider-location": "374", "resource-divider-location": "322",
        "version": "3.3.3308", "locale": "en",
    })
    ET.SubElement(root, "description")

    view = ET.SubElement(root, "view", {"zooming-state": "default:2", "id": "gantt-chart"})
    ET.SubElement(view, "field", {"id": "tpd3", "name": "Name", "width": "200", "order": "0"})
    ET.SubElement(view, "field", {"id": "tpd4", "name": "Begin date", "width": "75", "order": "1"})
    ET.SubElement(view, "field", {"id": "tpd5", "name": "End date", "width": "75", "order": "2"})
    ET.SubElement(view, "field", {"id": "tpd7", "name": "Completion", "width": "60", "order": "3"})

    resource_view = ET.SubElement(root, "view", {"id": "resource-table"})
    ET.SubElement(resource_view, "field", {"id": "0", "name": "Name", "width": "210", "order": "0"})

    taskprops = ET.SubElement(root, "taskproperties")
    for pid, pname, ptype, pvtype in [
        ("tpd0", "type", "default", "icon"), ("tpd1", "priority", "default", "icon"),
        ("tpd2", "info", "default", "icon"), ("tpd3", "name", "default", "text"),
        ("tpd4", "begindate", "default", "date"), ("tpd5", "enddate", "default", "date"),
        ("tpd6", "duration", "default", "int"), ("tpd7", "completion", "default", "int"),
        ("tpd8", "coordinator", "default", "text"), ("tpd9", "predecessorsr", "default", "text"),
    ]:
        ET.SubElement(taskprops, "taskproperty", {"id": pid, "name": pname, "type": ptype, "valuetype": pvtype})

    tasks_root = ET.SubElement(root, "tasks", {"empty-milestones": "true"})
    id_counter = [0]

    def next_id():
        v = id_counter[0]
        id_counter[0] += 1
        return v

    def completion(tid):
        return STATUS_COMPLETION.get(get_entry(state, tid)["status"], 0)

    def color(tid):
        return GAN_COLORS.get(get_entry(state, tid)["status"], GAN_COLORS["not-started"])

    for section_name, subsections in sections.items():
        sec_tids = [t for sub in subsections.values() for t in sub]
        sec_start = min(TASKS[t]["start"] for t in sec_tids)
        sec_end = max(leaf_end(state, t) for t in sec_tids)
        sec_duration = max(1, (sec_end - sec_start).days)
        sec_complete = round(sum(completion(t) for t in sec_tids) / len(sec_tids))
        sec_el = ET.SubElement(tasks_root, "task", {
            "id": str(next_id()), "name": section_name, "color": COLORS["blue"], "meeting": "false",
            "start": sec_start.isoformat(), "duration": str(sec_duration), "complete": str(sec_complete), "expand": "true",
        })

        for sub_name, tids in subsections.items():
            sub_start = min(TASKS[t]["start"] for t in tids)
            sub_end = max(leaf_end(state, t) for t in tids)
            sub_duration = max(1, (sub_end - sub_start).days)
            sub_complete = round(sum(completion(t) for t in tids) / len(tids))
            sub_el = ET.SubElement(sec_el, "task", {
                "id": str(next_id()), "name": sub_name, "color": COLORS["upcoming"], "meeting": "false",
                "start": sub_start.isoformat(), "duration": str(sub_duration), "complete": str(sub_complete), "expand": "true",
            })

            for tid in tids:
                task = TASKS[tid]
                end_date = leaf_end(state, tid)
                duration = max(1, (end_date - task["start"]).days)
                entry = get_entry(state, tid)
                name = f"{task['id']} {task['name']}"
                if entry.get("revised", "").strip():
                    name += "  [revised]"
                leaf_el = ET.SubElement(sub_el, "task", {
                    "id": str(next_id()), "name": name, "color": color(tid), "meeting": "false",
                    "start": task["start"].isoformat(), "duration": str(duration), "complete": str(completion(tid)), "expand": "true",
                })
                # Keep compatibility with legacy notes while also exporting current structured notes.
                note_lines = []
                if entry.get("notes"):
                    note_lines.append(entry["notes"])
                for n in entry.get("note_entries", []):
                    note_lines.append(f"[{n.get('date') or 'undated'}] {NOTE_CATEGORY_LABELS.get(n.get('category'), n.get('category'))}: {n.get('text','')}")
                if note_lines:
                    notes_el = ET.SubElement(leaf_el, "notes")
                    notes_el.text = "\n".join(note_lines)

    for tag in ["resources", "allocations", "vacations", "previous"]:
        ET.SubElement(root, tag)
    ET.SubElement(root, "roles", {"roleset-name": "Default"})

    rough = ET.tostring(root, encoding="utf-8")
    return minidom.parseString(rough).toprettyxml(indent="  ", encoding="UTF-8")


def sorted_filtered_ids(state, today, filter_key="all", sort_mode="Plan order"):
    ids = []
    for tid in TASK_ORDER:
        entry = get_entry(state, tid)
        bucket = bucket_for(TASKS[tid], entry, today)
        if filter_key != "all" and bucket != filter_key:
            continue
        ids.append(tid)

    if sort_mode == "Plan order":
        return ids
    if sort_mode == "Due date (soonest first)":
        return sorted(ids, key=lambda t: TASKS[t]["end"])
    if sort_mode == "Most overdue first":
        return sorted(ids, key=lambda t: (TASKS[t]["end"] - today).days)
    if sort_mode == "Name (A–Z)":
        return sorted(ids, key=lambda t: TASKS[t]["name"].lower())
    if sort_mode == "Status":
        return sorted(ids, key=lambda t: STATUSES.index(get_entry(state, t)["status"]))
    if sort_mode == "Owner (A–Z)":
        def owner_key(t):
            owner = get_entry(state, t).get("owner", "")
            return (1, "", TASKS[t]["name"].lower()) if not owner else (0, owner, TASKS[t]["name"].lower())
        return sorted(ids, key=owner_key)
    return ids
